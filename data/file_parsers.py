import csv
import io
import os
import logging

logger = logging.getLogger(__name__)

SUPPORTED_EXTENSIONS = {'.csv', '.xlsx', '.pdf', '.docx', '.txt', '.rtf'}
TABULAR_EXTENSIONS = {'.csv', '.xlsx'}
DOCUMENT_EXTENSIONS = {'.pdf', '.docx', '.txt', '.rtf'}
MAX_DOCUMENT_WORDS = 5000
MAX_ROWS = 100
MAX_PREVIEW_ROWS = 20
MAX_HEADERS = 15


def _truncate_text(text, max_words=MAX_DOCUMENT_WORDS):
    words = text.split()
    if len(words) <= max_words:
        return text
    return ' '.join(words[:max_words]) + '\n\n[Document truncated — showing first 5,000 words]'


def parse_csv(file_bytes):
    try:
        content = file_bytes.decode('utf-8')
    except UnicodeDecodeError:
        content = file_bytes.decode('latin-1')
    reader = csv.DictReader(io.StringIO(content))
    rows = list(reader)

    if len(rows) > MAX_ROWS:
        rows = rows[:MAX_ROWS]

    headers = reader.fieldnames or []
    preview = [dict(row) for row in rows[:MAX_PREVIEW_ROWS]]

    return {
        'type': 'tabular',
        'headers': headers[:MAX_HEADERS],
        'rows': rows,
        'preview': preview,
        'row_count': len(rows),
        'summary': f"CSV file with {len(rows)} rows and columns: {', '.join(headers[:10])}"
    }


def parse_excel(file_bytes, filename):
    from openpyxl import load_workbook

    wb = load_workbook(io.BytesIO(file_bytes), read_only=True, data_only=True)
    ws = wb.active

    all_rows = []
    headers = []
    for i, row in enumerate(ws.iter_rows(values_only=True)):
        str_row = [str(cell) if cell is not None else '' for cell in row]
        if i == 0:
            headers = str_row
        else:
            all_rows.append(str_row)
        if i >= MAX_ROWS:
            break

    wb.close()

    rows_as_dicts = []
    for row_vals in all_rows:
        row_dict = {}
        for j, header in enumerate(headers):
            if j < len(row_vals):
                row_dict[header] = row_vals[j]
        rows_as_dicts.append(row_dict)

    preview = rows_as_dicts[:MAX_PREVIEW_ROWS]

    return {
        'type': 'tabular',
        'headers': headers[:MAX_HEADERS],
        'rows': rows_as_dicts,
        'preview': preview,
        'row_count': len(rows_as_dicts),
        'summary': f"Excel file with {len(rows_as_dicts)} rows and columns: {', '.join(headers[:10])}"
    }


MAX_PDF_PAGES = 50


def parse_pdf(file_bytes):
    from PyPDF2 import PdfReader

    reader = PdfReader(io.BytesIO(file_bytes))
    total_pages = len(reader.pages)
    pages = []
    for page in reader.pages[:MAX_PDF_PAGES]:
        text = page.extract_text()
        if text:
            pages.append(text.strip())

    full_text = '\n\n'.join(pages)
    truncated = _truncate_text(full_text)
    word_count = len(full_text.split())

    truncation_note = f" (first {MAX_PDF_PAGES} of {total_pages} pages)" if total_pages > MAX_PDF_PAGES else ""
    return {
        'type': 'document',
        'format': 'pdf',
        'text': truncated,
        'page_count': total_pages,
        'word_count': word_count,
        'summary': f"PDF document with {total_pages} pages{truncation_note} and ~{word_count} words"
    }


def parse_docx(file_bytes):
    from docx import Document

    doc = Document(io.BytesIO(file_bytes))
    paragraphs = []
    for para in doc.paragraphs:
        text = para.text.strip()
        if text:
            paragraphs.append(text)

    for table in doc.tables:
        table_rows = []
        for row in table.rows:
            cells = [cell.text.strip() for cell in row.cells]
            table_rows.append(' | '.join(cells))
        if table_rows:
            paragraphs.append('\n'.join(table_rows))

    full_text = '\n\n'.join(paragraphs)
    truncated = _truncate_text(full_text)
    word_count = len(full_text.split())

    return {
        'type': 'document',
        'format': 'docx',
        'text': truncated,
        'paragraph_count': len(paragraphs),
        'word_count': word_count,
        'summary': f"Word document with {len(paragraphs)} sections and ~{word_count} words"
    }


def parse_txt(file_bytes):
    try:
        content = file_bytes.decode('utf-8')
    except UnicodeDecodeError:
        content = file_bytes.decode('latin-1')

    lines = content.strip().split('\n')
    truncated = _truncate_text(content)
    word_count = len(content.split())

    return {
        'type': 'document',
        'format': 'txt',
        'text': truncated,
        'line_count': len(lines),
        'word_count': word_count,
        'summary': f"Text file with {len(lines)} lines and ~{word_count} words"
    }


def parse_rtf(file_bytes):
    from striprtf.striprtf import rtf_to_text

    try:
        rtf_content = file_bytes.decode('utf-8')
    except UnicodeDecodeError:
        rtf_content = file_bytes.decode('latin-1')

    plain_text = rtf_to_text(rtf_content)
    truncated = _truncate_text(plain_text)
    word_count = len(plain_text.split())

    return {
        'type': 'document',
        'format': 'rtf',
        'text': truncated,
        'word_count': word_count,
        'summary': f"RTF document with ~{word_count} words"
    }


def parse_file(file_bytes, filename):
    ext = os.path.splitext(filename.lower())[1]

    if ext not in SUPPORTED_EXTENSIONS:
        raise ValueError(f"Unsupported file format: {ext}. Supported formats: {', '.join(sorted(SUPPORTED_EXTENSIONS))}")

    if ext == '.csv':
        return parse_csv(file_bytes)
    elif ext == '.xlsx':
        return parse_excel(file_bytes, filename)
    elif ext == '.pdf':
        return parse_pdf(file_bytes)
    elif ext == '.docx':
        return parse_docx(file_bytes)
    elif ext == '.txt':
        return parse_txt(file_bytes)
    elif ext == '.rtf':
        return parse_rtf(file_bytes)
    else:
        raise ValueError(f"Unsupported file format: {ext}")
